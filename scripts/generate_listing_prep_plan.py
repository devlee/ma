# -*- coding: utf-8 -*-
"""
根据月度上架计划，反推上架准备工作(选款/打分/定款/SPU/作图/同步)的按周排期，
输出 Excel 计划表。

规则要点：
- 复杂品类(妈妈装MBD/婚纱WD, 60%+单个作图): 上架前5周启动选款, 作图3周
- 常规品类(约30%单个作图): 上架前4周启动选款, 作图2周
- 小批次(≤20款): 上架前3周合并"选款+打分"(可直接定款), 作图2周
- 每周上新: ≥40款按月内各周均摊; 21-39款分2周; ≤20款集中1周
- 春节: 2027/2/5-2/18 放假, 2/8与2/15两周不排工作, 2月批次全部节前完成作图
- 8月为追赶月, 按当前实际进度单独排期
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- 基础数据

CATS = ["MBD", "ED", "WD", "BD", "Prom", "CD", "FG", "JBD"]
CAT_NAMES = {
    "MBD": "妈妈装 Mother of the Bride",
    "ED": "晚礼服 Evening",
    "WD": "婚纱 Bride",
    "BD": "伴娘裙 Bridesmaid",
    "Prom": "舞会裙 Prom",
    "CD": "鸡尾酒裙 Cocktail",
    "FG": "花童裙 Flower Girl",
    "JBD": "小伴娘裙 JBD",
}
COMPLEX_CATS = {"MBD", "WD"}          # 60%以上款需单个作图
SINGLE_RATIO = {c: (0.6 if c in COMPLEX_CATS else 0.3) for c in CATS}

MONTHS = ["2026-08", "2026-09", "2026-10", "2026-11", "2026-12", "2027-01", "2027-02"]
PLAN = {  # 月度上架计划(款)
    "MBD":  [70, 80, 80, 100, 80, 80, 50],
    "ED":   [40, 40, 40, 60, 50, 40, 30],
    "WD":   [50, 50, 40, 40, 10, 10, 10],
    "BD":   [40, 40, 40, 50, 40, 40, 30],
    "Prom": [0, 0, 0, 20, 40, 40, 0],
    "CD":   [0, 20, 10, 20, 20, 30, 0],
    "FG":   [0, 30, 20, 10, 20, 20, 0],
    "JBD":  [0, 0, 30, 0, 20, 0, 0],
}

# 周(周一日期)序列: 2026-08-03 ~ 2027-03-01
FIRST_MONDAY = dt.date(2026, 8, 3)
WEEKS = [FIRST_MONDAY + dt.timedelta(weeks=i) for i in range(31)]
HOLIDAY_WEEKS = {dt.date(2027, 2, 8), dt.date(2027, 2, 15)}  # 春节 2/5-2/18

def W(y, m, d):
    return dt.date(y, m, d)

def wk_label(monday):
    fri = monday + dt.timedelta(days=4)
    return f"{monday.month}/{monday.day}周\n({monday.month}.{monday.day}-{fri.month}.{fri.day})"

# 各月上架周(周一)
SYNC_WEEKS = {
    "2026-08": [W(2026, 8, 10), W(2026, 8, 17), W(2026, 8, 24), W(2026, 8, 31)],
    "2026-09": [W(2026, 9, 7), W(2026, 9, 14), W(2026, 9, 21), W(2026, 9, 28)],
    "2026-10": [W(2026, 10, 5), W(2026, 10, 12), W(2026, 10, 19), W(2026, 10, 26)],
    "2026-11": [W(2026, 11, 2), W(2026, 11, 9), W(2026, 11, 16), W(2026, 11, 23), W(2026, 11, 30)],
    "2026-12": [W(2026, 12, 7), W(2026, 12, 14), W(2026, 12, 21), W(2026, 12, 28)],
    "2027-01": [W(2027, 1, 4), W(2027, 1, 11), W(2027, 1, 18), W(2027, 1, 25)],
    "2027-02": [W(2027, 2, 1), W(2027, 2, 22)],  # 春节: 2/8、2/15两周放假
}

# ------------------------------------------------------- 批次定义(含8月追赶)
# 每个批次: cat, month, qty, sel/score/ding/spu/art 各阶段所在周, sync {周:款数}, note

def weeks_before(monday, n):
    return monday - dt.timedelta(weeks=n)

def even_split(qty, weeks):
    n = len(weeks)
    base, rem = divmod(qty, n)
    return {w: base + (1 if i < rem else 0) for i, w in enumerate(weeks)}

BATCHES = []

def add_batch(cat, month, qty, sel, score, ding, spu, art, sync, note=""):
    assert sum(sync.values()) == qty, f"{cat} {month} 同步数与批次量不符"
    BATCHES.append(dict(cat=cat, month=month, qty=qty, sel=sel, score=score,
                        ding=ding, spu=spu, art=art, sync=sync, note=note))

# ---- 2026-08 追赶批次(按当前实际进度) ----
add_batch("MBD", "2026-08", 70,
          sel=[], score=[], ding=[], spu=[],
          art=[W(2026, 8, 3), W(2026, 8, 10)],
          sync={W(2026, 8, 10): 20, W(2026, 8, 17): 20, W(2026, 8, 24): 15, W(2026, 8, 31): 15},
          note="已在作图环节, 边作图边滚动上架")
add_batch("ED", "2026-08", 40,
          sel=[], score=[], ding=[], spu=[W(2026, 8, 3)],
          art=[W(2026, 8, 3), W(2026, 8, 10)],
          sync={W(2026, 8, 17): 15, W(2026, 8, 24): 15, W(2026, 8, 31): 10},
          note="定款已完成, 本周生成SPU并启动作图")
add_batch("BD", "2026-08", 40,
          sel=[], score=[], ding=[W(2026, 8, 3)], spu=[W(2026, 8, 10)],
          art=[W(2026, 8, 10), W(2026, 8, 17)],
          sync={W(2026, 8, 24): 20, W(2026, 8, 31): 20},
          note="选款打分已完成, 本周内定款")
add_batch("WD", "2026-08", 50,
          sel=[W(2026, 8, 3)], score=[], ding=[W(2026, 8, 3)], spu=[W(2026, 8, 10)],
          art=[W(2026, 8, 10), W(2026, 8, 17), W(2026, 8, 24)],
          sync={W(2026, 8, 24): 15, W(2026, 8, 31): 35},
          note="尚未启动, 高风险: 建议本周直接定款(跳过打分), 作图压缩; 如仍延期可将部分顺延至9/7周")

# ---- 2026-09 起: 标准节奏反推 ----
def std_schedule(cat, month, qty):
    m_idx = MONTHS.index(month)
    sync_wks = SYNC_WEEKS[month]
    w0 = sync_wks[0]
    small = qty <= 20
    if small:
        # 小批次: 集中1周上架(周次在月内错开, 平衡周上架量)
        pass
    # 同步拆分
    if qty >= 40:
        sync = even_split(qty, sync_wks)
    elif qty > 20:
        two = [sync_wks[0], sync_wks[2 if len(sync_wks) > 2 else -1]]
        sync = even_split(qty, two)
    else:
        sync = None  # 由调用方指定集中上架周
    return w0, small, sync

# 小批次(≤20款)集中上架周: 手工错开以平衡每周总量
SMALL_SYNC_WEEK = {
    ("CD", "2026-09"): W(2026, 9, 14),
    ("CD", "2026-10"): W(2026, 10, 12),
    ("FG", "2026-10"): W(2026, 10, 26),
    ("Prom", "2026-11"): W(2026, 11, 9),
    ("CD", "2026-11"): W(2026, 11, 16),
    ("FG", "2026-11"): W(2026, 11, 23),
    ("CD", "2026-12"): W(2026, 12, 7),
    ("WD", "2026-12"): W(2026, 12, 14),
    ("FG", "2026-12"): W(2026, 12, 21),
    ("JBD", "2026-12"): W(2026, 12, 28),
    ("WD", "2027-01"): W(2027, 1, 11),
    ("FG", "2027-01"): W(2027, 1, 25),
    ("WD", "2027-02"): W(2027, 2, 1),
}

for month in MONTHS[1:]:
    m_idx = MONTHS.index(month)
    for cat in CATS:
        qty = PLAN[cat][m_idx]
        if qty == 0:
            continue
        w0, small, sync = std_schedule(cat, month, qty)
        if small:
            sw = SMALL_SYNC_WEEK[(cat, month)]
            sync = {sw: qty}
            # 小批次: 前3周合并选款+打分(可直接定款), 前2周SPU+作图
            add_batch(cat, month, qty,
                      sel=[weeks_before(sw, 3)], score=[weeks_before(sw, 3)],
                      ding=[weeks_before(sw, 3)], spu=[weeks_before(sw, 2)],
                      art=[weeks_before(sw, 2), weeks_before(sw, 1)],
                      sync=sync,
                      note="小批次: 选款打分同周合并, 可走直接定款再压缩1周")
        elif cat in COMPLEX_CATS:
            # 复杂品类: 前5周选款, 前4周打分+定款, 前3周SPU, 作图3周
            add_batch(cat, month, qty,
                      sel=[weeks_before(w0, 5)], score=[weeks_before(w0, 4)],
                      ding=[weeks_before(w0, 4)], spu=[weeks_before(w0, 3)],
                      art=[weeks_before(w0, 3), weeks_before(w0, 2), weeks_before(w0, 1)],
                      sync=sync,
                      note="复杂品类(60%+单个作图), 作图排3周")
        else:
            # 常规品类: 前4周选款, 前3周打分+定款, 前2周SPU, 作图2周
            add_batch(cat, month, qty,
                      sel=[weeks_before(w0, 4)], score=[weeks_before(w0, 3)],
                      ding=[weeks_before(w0, 3)], spu=[weeks_before(w0, 2)],
                      art=[weeks_before(w0, 2), weeks_before(w0, 1)],
                      sync=sync,
                      note="")

# 校验: 各月合计
for m_idx, month in enumerate(MONTHS):
    total = sum(b["qty"] for b in BATCHES if b["month"] == month)
    expect = sum(PLAN[c][m_idx] for c in CATS)
    assert total == expect, f"{month}: {total} != {expect}"

# ---------------------------------------------------------------- 样式

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="4472C4")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
BASE_FONT = Font(size=10)
BOLD = Font(bold=True, size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

STAGE_FILL = {
    "选款": PatternFill("solid", fgColor="BDD7EE"),
    "打分": PatternFill("solid", fgColor="FFE699"),
    "定款": PatternFill("solid", fgColor="F8CBAD"),
    "SPU": PatternFill("solid", fgColor="D9D2E9"),
    "作图": PatternFill("solid", fgColor="C9DAF8"),
    "上架": PatternFill("solid", fgColor="C6EFCE"),
    "假期": PatternFill("solid", fgColor="E7E6E6"),
}

def style_header(cell):
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = CENTER
    cell.border = BORDER

def style_cell(cell, align=CENTER, font=BASE_FONT):
    cell.alignment = align
    cell.border = BORDER
    cell.font = font

# ---------------------------------------------------------------- 生成工作簿

wb = Workbook()

# ===== Sheet 1: 说明与规则 =====
ws = wb.active
ws.title = "说明与规则"
rules = [
    ("目的", "根据 2026-08 ~ 2027-02 月度上架计划, 反推选款/打分/定款/生成SPU/作图/同步(上架)各环节的按周节点安排, 保证每月上架量按周均匀完成。"),
    ("周定义", "以周一日期标识一周(工作周为周一至周五)。"),
    ("环节耗时", "选款 3-5个工作日(量大取5天); 打分 2-3个工作日, 打分完成当周内定款; 生成SPU 2-3个工作日; 作图: 常规品类约2周(约30%款单个制作), 妈妈装/婚纱约3周(60%以上款单个制作)。"),
    ("反推提前量", "妈妈装/婚纱: 上架前第5周启动选款; 常规品类: 前第4周; 小批次(≤20款): 前第3周合并选款+打分(或直接定款), 可再压缩。"),
    ("上架节奏", "每周上新: 月量≥40款的品类按月内各周均摊; 21-39款分2周; ≤20款集中1周上架(各品类错开周次, 平衡周上架总量)。"),
    ("选款产能", "选款池按定款量的1.5倍备选(可调); 人均选款产能: 大批次(≥40款)一人每周40款, 小批次一人每周20款。详见\"产能与作图拆解\"表。"),
    ("直接定款(灵活机制)", "若某周选款/打分排不开, 小批次品类(CD/FG/JBD/Prom及低量月份的WD)可由负责人直接定款, 跳过组内打分, 整体可压缩1-2周。甘特为标准节奏, 实际执行允许±1周浮动, 但作图完成时间不得晚于对应上架周的前一个周五。"),
    ("春节安排", "假设 2027/2/5-2/18 放假14天(以公司通知为准), 2/8周与2/15周不排任何工作。2月120款全部于 1/29(节前)完成作图, 分 2/1周(65款)与 2/22周(55款)两次上架。"),
    ("2027-03 批次提示", "受春节影响, 3月批次(量待定)中妈妈装/婚纱需提前: 1/18-1/29 完成选款打分定款, 节前完成SPU并启动作图, 节后 2/22周 收尾作图, 3月第1周正常上架; 常规品类节后 2/22周 立即选款(或直接定款)。"),
    ("8月追赶(当前进度)", "BD: 选款打分已完成→本周定款; ED: 已定款→本周SPU+作图; MBD: 作图中→8/10周起滚动上架; WD: 未启动→高风险, 建议本周直接定款并压缩作图, 8/24周起上架, 必要时部分顺延至9月第1周。8月上架量后置明显(8/24、8/31两周合计145款), 属追赶期特殊情况。"),
    ("图例", "选款=浅蓝, 打分=黄色, 定款=橙色, 生成SPU=浅紫, 作图=蓝紫, 上架=绿色, 春节假期=灰色。"),
]
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 120
c = ws.cell(row=1, column=1, value="上架准备工作按周排期 — 说明与规则 (2026-08 ~ 2027-02)")
c.font = Font(bold=True, size=13)
r = 3
for k, v in rules:
    kc = ws.cell(row=r, column=1, value=k)
    vc = ws.cell(row=r, column=2, value=v)
    style_cell(kc, LEFT, BOLD)
    style_cell(vc, LEFT)
    ws.row_dimensions[r].height = max(28, 14 * (len(v) // 60 + 1))
    r += 1

# ===== Sheet 2: 月度上架计划 =====
ws = wb.create_sheet("月度上架计划")
ws.cell(row=1, column=1, value="品类")
for j, m in enumerate(MONTHS):
    ws.cell(row=1, column=2 + j, value=m)
ws.cell(row=1, column=2 + len(MONTHS), value="合计")
for cell in ws[1]:
    if cell.value:
        style_header(cell)
for i, cat in enumerate(CATS):
    r = 2 + i
    c = ws.cell(row=r, column=1, value=CAT_NAMES[cat])
    style_cell(c, LEFT, BOLD)
    for j in range(len(MONTHS)):
        cc = ws.cell(row=r, column=2 + j, value=PLAN[cat][j] or "")
        style_cell(cc)
    tc = ws.cell(row=r, column=2 + len(MONTHS), value=sum(PLAN[cat]))
    style_cell(tc, CENTER, BOLD)
r = 2 + len(CATS)
c = ws.cell(row=r, column=1, value="月合计")
style_cell(c, LEFT, BOLD)
for j in range(len(MONTHS)):
    cc = ws.cell(row=r, column=2 + j, value=sum(PLAN[c2][j] for c2 in CATS))
    style_cell(cc, CENTER, BOLD)
tc = ws.cell(row=r, column=2 + len(MONTHS), value=sum(sum(PLAN[c2]) for c2 in CATS))
style_cell(tc, CENTER, BOLD)
ws.column_dimensions["A"].width = 28
for j in range(len(MONTHS) + 1):
    ws.column_dimensions[get_column_letter(2 + j)].width = 10

# ===== Sheet 3: 每周上架拆解 =====
ws = wb.create_sheet("每周上架拆解")
sync_week_list = [w for w in WEEKS if any(w in b["sync"] for b in BATCHES) or w in HOLIDAY_WEEKS]
ws.cell(row=1, column=1, value="品类")
for j, w in enumerate(sync_week_list):
    hc = ws.cell(row=1, column=2 + j, value=wk_label(w))
    style_header(hc)
style_header(ws.cell(row=1, column=1))
tc = ws.cell(row=1, column=2 + len(sync_week_list), value="合计")
style_header(tc)
for i, cat in enumerate(CATS):
    r = 2 + i
    c = ws.cell(row=r, column=1, value=CAT_NAMES[cat])
    style_cell(c, LEFT, BOLD)
    row_total = 0
    for j, w in enumerate(sync_week_list):
        qty = sum(b["sync"].get(w, 0) for b in BATCHES if b["cat"] == cat)
        row_total += qty
        cc = ws.cell(row=r, column=2 + j, value=qty or "")
        style_cell(cc)
        if w in HOLIDAY_WEEKS:
            cc.fill = STAGE_FILL["假期"]
        elif qty:
            cc.fill = STAGE_FILL["上架"]
    tc = ws.cell(row=r, column=2 + len(sync_week_list), value=row_total)
    style_cell(tc, CENTER, BOLD)
r = 2 + len(CATS)
c = ws.cell(row=r, column=1, value="周合计")
style_cell(c, LEFT, BOLD)
grand = 0
for j, w in enumerate(sync_week_list):
    qty = sum(b["sync"].get(w, 0) for b in BATCHES)
    grand += qty
    cc = ws.cell(row=r, column=2 + j, value=qty if w not in HOLIDAY_WEEKS else "春节")
    style_cell(cc, CENTER, BOLD)
    if w in HOLIDAY_WEEKS:
        cc.fill = STAGE_FILL["假期"]
tc = ws.cell(row=r, column=2 + len(sync_week_list), value=grand)
style_cell(tc, CENTER, BOLD)
ws.column_dimensions["A"].width = 28
for j in range(len(sync_week_list) + 1):
    ws.column_dimensions[get_column_letter(2 + j)].width = 11
ws.freeze_panes = "B2"

# ===== Sheet 4: 批次甘特总览 =====
ws = wb.create_sheet("批次甘特总览")
headers = ["品类", "月份批次", "款数"] + [wk_label(w) for w in WEEKS] + ["备注"]
for j, h in enumerate(headers):
    hc = ws.cell(row=1, column=1 + j, value=h)
    style_header(hc)
STAGE_ORDER = [("sel", "选款"), ("score", "打分"), ("ding", "定款"), ("spu", "SPU"), ("art", "作图")]
r = 2
for month in MONTHS:
    for cat in CATS:
        bs = [b for b in BATCHES if b["month"] == month and b["cat"] == cat]
        if not bs:
            continue
        b = bs[0]
        style_cell(ws.cell(row=r, column=1, value=CAT_NAMES[cat]), LEFT, BOLD)
        style_cell(ws.cell(row=r, column=2, value=f"{month} 批"), CENTER, BOLD)
        style_cell(ws.cell(row=r, column=3, value=b["qty"]))
        for j, w in enumerate(WEEKS):
            parts = []
            fill_key = None
            for key, label in STAGE_ORDER:
                if w in b[key]:
                    lb = label
                    if key == "score":
                        lb = "打分"
                    parts.append(lb)
                    fill_key = {"sel": "选款", "score": "打分", "ding": "定款", "spu": "SPU", "art": "作图"}[key]
            if w in b["sync"]:
                parts.append(f"上架{b['sync'][w]}")
                fill_key = "上架"
            cc = ws.cell(row=r, column=4 + j, value="+".join(parts) if parts else "")
            style_cell(cc)
            if w in HOLIDAY_WEEKS:
                cc.fill = STAGE_FILL["假期"]
                if not parts:
                    cc.value = "假期"
            elif fill_key:
                cc.fill = STAGE_FILL[fill_key]
        nc = ws.cell(row=r, column=4 + len(WEEKS), value=b["note"])
        style_cell(nc, LEFT)
        r += 1
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 11
ws.column_dimensions["C"].width = 6
for j in range(len(WEEKS)):
    ws.column_dimensions[get_column_letter(4 + j)].width = 10.5
ws.column_dimensions[get_column_letter(4 + len(WEEKS))].width = 46
ws.freeze_panes = "D2"

# ===== Sheet 5: 每周工作清单 =====
ws = wb.create_sheet("每周工作清单")
cols = ["周次", "选款", "打分+定款", "生成SPU", "作图(进行中)", "同步上架", "本周上架量", "备注"]
for j, h in enumerate(cols):
    style_header(ws.cell(row=1, column=1 + j, value=h))

def batch_tag(b):
    return f"{CAT_NAMES[b['cat']].split(' ')[0]}{b['month'][2:].replace('-', '')}批({b['qty']}款)"

r = 2
for w in WEEKS:
    style_cell(ws.cell(row=r, column=1, value=wk_label(w).replace("\n", " ")), LEFT, BOLD)
    if w in HOLIDAY_WEEKS:
        cc = ws.cell(row=r, column=2, value="春节假期")
        style_cell(cc, CENTER, BOLD)
        cc.fill = STAGE_FILL["假期"]
        for j in range(3, len(cols) + 1):
            c2 = ws.cell(row=r, column=j, value="")
            style_cell(c2)
            c2.fill = STAGE_FILL["假期"]
        r += 1
        continue
    sel = [batch_tag(b) for b in BATCHES if w in b["sel"]]
    score = [batch_tag(b) for b in BATCHES if w in b["score"] or w in b["ding"]]
    spu = [batch_tag(b) for b in BATCHES if w in b["spu"]]
    art = [batch_tag(b) for b in BATCHES if w in b["art"]]
    sync = [f"{CAT_NAMES[b['cat']].split(' ')[0]} {b['sync'][w]}款" for b in BATCHES if w in b["sync"]]
    total = sum(b["sync"][w] for b in BATCHES if w in b["sync"])
    notes = []
    if w == W(2026, 8, 3):
        notes.append("追赶周: WD 8月批紧急直接定款; 同时启动 MBD/WD 9月批选款")
    if w == W(2027, 1, 25):
        notes.append("节前最后完整周: 确认2月批次作图全部完成; 启动3月批次(MBD/WD)选款")
    if w == W(2027, 2, 1):
        notes.append("节前上架周(2/5起放假)")
    if w == W(2027, 2, 22):
        notes.append("节后复工: 上架2月剩余批次; 3月批次常规品类立即选款/定款")
    vals = ["、".join(sel), "、".join(score), "、".join(spu), "、".join(art),
            "、".join(sync), total or "", "; ".join(notes)]
    for j, v in enumerate(vals):
        cc = ws.cell(row=r, column=2 + j, value=v)
        style_cell(cc, LEFT if j != 5 else CENTER)
    ws.row_dimensions[r].height = 30
    r += 1
widths = [16, 30, 30, 26, 34, 34, 10, 46]
for j, wd_ in enumerate(widths):
    ws.column_dimensions[get_column_letter(1 + j)].width = wd_
ws.freeze_panes = "B2"

# ===== Sheet 6: 产能与作图拆解 =====
ws = wb.create_sheet("产能与作图拆解")
cols = ["月份批次", "品类", "定款量", "选款池(×1.5)", "人均选款产能(款/周)", "选款需人·周",
        "单个制作占比", "单个制作(款)", "批量制作(款)", "作图周期(周)"]
for j, h in enumerate(cols):
    style_header(ws.cell(row=1, column=1 + j, value=h))
r = 2
import math
for month in MONTHS:
    for cat in CATS:
        bs = [b for b in BATCHES if b["month"] == month and b["cat"] == cat]
        if not bs:
            continue
        b = bs[0]
        qty = b["qty"]
        pool = math.ceil(qty * 1.5)
        cap = 40 if qty >= 40 else 20
        ppl = round(pool / cap, 1)
        single = round(qty * SINGLE_RATIO[cat])
        art_weeks = len(b["art"]) if b["art"] else ""
        vals = [f"{month} 批", CAT_NAMES[cat], qty, pool, cap, ppl,
                f"{int(SINGLE_RATIO[cat]*100)}%", single, qty - single, art_weeks]
        for j, v in enumerate(vals):
            cc = ws.cell(row=r, column=1 + j, value=v)
            style_cell(cc, LEFT if j == 1 else CENTER)
        r += 1
widths2 = [12, 26, 9, 12, 18, 12, 12, 12, 12, 12]
for j, wd_ in enumerate(widths2):
    ws.column_dimensions[get_column_letter(1 + j)].width = wd_
ws.freeze_panes = "A2"

OUT = "/workspace/上架准备工作按周排期_2026-08_2027-02.xlsx"
wb.save(OUT)
print("saved:", OUT)

# 汇总校验输出
print("\n各月上架合计(按周拆解求和):")
for m_idx, month in enumerate(MONTHS):
    total = sum(sum(b["sync"].values()) for b in BATCHES if b["month"] == month)
    print(f"  {month}: {total} (计划 {sum(PLAN[c][m_idx] for c in CATS)})")
print("\n每周上架总量:")
for w in WEEKS:
    t = sum(b["sync"].get(w, 0) for b in BATCHES)
    if t:
        print(f"  {w}: {t}")
